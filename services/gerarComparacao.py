import pandas as pd
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def gerar_excel_comparacao(notas: list[dict]) -> str:
    """
    Gera Excel comparando preços de um mesmo item entre diferentes NFC-es.
    notas: lista de dicts com {"rotulo": str, "itens": list[dict]}
    Retorna o caminho do arquivo.
    """
    caminho = "comparacao_precos.xlsx"

    azul_escuro = "1E3A5F"
    azul_medio  = "1E6FD9"
    cinza_linha = "F2F6FC"
    branco      = "FFFFFF"
    verde       = "C6EFCE"
    verde_fonte = "276221"
    vermelho    = "FFC7CE"
    verm_fonte  = "9C0006"

    borda_fina = Border(
        left=Side(style="thin", color="B0C4DE"),
        right=Side(style="thin", color="B0C4DE"),
        top=Side(style="thin", color="B0C4DE"),
        bottom=Side(style="thin", color="B0C4DE"),
    )

    # Monta dicionário: produto → {rotulo: preço}
    mapa: dict[str, dict[str, float]] = {}
    rotulos = [n["rotulo"] for n in notas]

    for nota in notas:
        for item in nota["itens"]:
            nome = item["Produto"].upper().strip()
            if nome not in mapa:
                mapa[nome] = {}
            mapa[nome][nota["rotulo"]] = item["Valor Unitário"]

    with pd.ExcelWriter(caminho, engine="openpyxl") as writer:
        # ── Aba Comparativo ──────────────────────────────────────────
        ws = writer.book.create_sheet("Comparativo", 0)

        n_cols = 1 + len(rotulos)  # Produto + N lojas
        ultima_col = get_column_letter(n_cols + 1)  # +1 para col Menor Preço

        ws.merge_cells(f"A1:{ultima_col}2")
        ws["A1"].value = "💰  COMPARATIVO DE PREÇOS"
        ws["A1"].font = Font(bold=True, size=14, color=branco)
        ws["A1"].fill = PatternFill("solid", fgColor=azul_escuro)
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 20
        ws.row_dimensions[2].height = 18

        # Cabeçalho da tabela
        cabecalhos = ["Produto"] + rotulos + ["✅ Menor Preço"]
        for col_idx, cab in enumerate(cabecalhos, start=1):
            c = ws.cell(row=4, column=col_idx, value=cab)
            c.font = Font(bold=True, color=branco, size=11)
            c.fill = PatternFill("solid", fgColor=azul_medio)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[4].height = 28

        # Linhas de dados
        for row_idx, (produto, precos) in enumerate(sorted(mapa.items()), start=5):
            fill_base = PatternFill("solid", fgColor=cinza_linha if row_idx % 2 == 0 else branco)

            c = ws.cell(row=row_idx, column=1, value=produto.title())
            c.fill = fill_base
            c.border = borda_fina

            valores_disponiveis = [precos.get(r) for r in rotulos]
            valores_numericos = [v for v in valores_disponiveis if v is not None]
            menor = min(valores_numericos) if valores_numericos else None

            for col_idx, rotulo in enumerate(rotulos, start=2):
                valor = precos.get(rotulo)
                c = ws.cell(row=row_idx, column=col_idx, value=valor if valor is not None else "—")
                c.border = borda_fina
                c.alignment = Alignment(horizontal="center")
                if valor is not None:
                    c.number_format = 'R$ #,##0.00'
                    if menor is not None and valor == menor and len(valores_numericos) > 1:
                        c.fill = PatternFill("solid", fgColor=verde)
                        c.font = Font(bold=True, color=verde_fonte)
                    elif menor is not None and len(valores_numericos) > 1:
                        c.fill = PatternFill("solid", fgColor=vermelho)
                        c.font = Font(color=verm_fonte)
                    else:
                        c.fill = fill_base
                else:
                    c.fill = fill_base

            # Col menor preço
            col_menor = len(rotulos) + 2
            c_menor = ws.cell(row=row_idx, column=col_menor)
            if menor is not None:
                # Encontra qual loja tem o menor preço
                loja_menor = [r for r in rotulos if precos.get(r) == menor]
                c_menor.value = f"{loja_menor[0]} — R$ {menor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
            else:
                c_menor.value = "—"
            c_menor.fill = fill_base
            c_menor.border = borda_fina
            c_menor.alignment = Alignment(horizontal="center")
            ws.row_dimensions[row_idx].height = 16

        # Largura das colunas
        ws.column_dimensions["A"].width = 40
        for col_idx in range(2, len(rotulos) + 3):
            ws.column_dimensions[get_column_letter(col_idx)].width = 20

        # ── Aba Dados Brutos ─────────────────────────────────────────
        for nota in notas:
            df = pd.DataFrame(nota["itens"])
            sheet_name = nota["rotulo"][:31]  # Limite do Excel
            df.to_excel(writer, index=False, sheet_name=sheet_name)

    return caminho
