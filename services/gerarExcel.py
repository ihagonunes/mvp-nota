import pandas as pd
import os
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def gerar_excel(itens: list[dict], pessoas: list[str]) -> str:
    """
    Gera Excel de divisão de compras com colunas corretas e aba Resumo.
    Retorna o caminho do arquivo.
    """
    caminho = "divisao_compras.xlsx"

    # ── Paleta ───────────────────────────────────────────────────────
    azul_escuro = "1E3A5F"
    azul_medio  = "1E6FD9"
    azul_claro  = "D6E4F7"
    cinza_linha = "F2F6FC"
    branco      = "FFFFFF"
    verde       = "1A7A4A"

    borda_fina = Border(
        left=Side(style="thin", color="B0C4DE"),
        right=Side(style="thin", color="B0C4DE"),
        top=Side(style="thin", color="B0C4DE"),
        bottom=Side(style="thin", color="B0C4DE"),
    )

    def aplicar_borda_range(ws, min_row, max_row, min_col, max_col):
        for row in ws.iter_rows(
            min_row=min_row, max_row=max_row,
            min_col=min_col, max_col=max_col,
        ):
            for cell in row:
                cell.border = borda_fina

    n_pessoas = len(pessoas) if pessoas else 1

    # Enriquece itens com colunas de divisão
    itens_enriquecidos = []
    for item in itens:
        subtotal = round(item["Quantidade"] * item["Valor Unitário"], 2)
        itens_enriquecidos.append({
            "Produto":           item["Produto"],
            "Quantidade":        item["Quantidade"],
            "Valor Unitário":    item["Valor Unitário"],
            "Subtotal":          subtotal,
            "Categoria":         "",          # preenchível pelo usuário
            "Nº de Pessoas":     n_pessoas,
            "Valor por Pessoa":  round(subtotal / n_pessoas, 2),
        })

    df = pd.DataFrame(itens_enriquecidos)
    total_geral = df["Subtotal"].sum()

    with pd.ExcelWriter(caminho, engine="openpyxl") as writer:
        # ── Aba Compras ──────────────────────────────────────────────
        df.to_excel(writer, index=False, sheet_name="Compras", startrow=1)
        ws = writer.book["Compras"]

        # Título da aba
        ws.merge_cells("A1:G1")
        ws["A1"].value = "🛒  DIVISÃO DE COMPRAS"
        ws["A1"].font = Font(bold=True, size=14, color=branco)
        ws["A1"].fill = PatternFill("solid", fgColor=azul_escuro)
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 24

        # Cabeçalho da tabela (linha 2)
        cabecalhos = [
            "Produto", "Quantidade", "Valor Unitário (R$)",
            "Subtotal (R$)", "Categoria", "Nº de Pessoas", "Valor por Pessoa (R$)"
        ]
        for col_idx, cab in enumerate(cabecalhos, start=1):
            c = ws.cell(row=2, column=col_idx, value=cab)
            c.font = Font(bold=True, color=branco, size=11)
            c.fill = PatternFill("solid", fgColor=azul_medio)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[2].height = 28

        # Linhas de dados — substituir fórmulas nas colunas D e G
        for i in range(3, len(itens_enriquecidos) + 3):
            fill = PatternFill("solid", fgColor=cinza_linha if i % 2 == 0 else branco)
            for col_idx in range(1, 8):
                ws.cell(row=i, column=col_idx).fill = fill

            # Col D: Subtotal = Qtd * Vl Unitário
            ws[f"D{i}"] = f"=B{i}*C{i}"
            ws[f"D{i}"].number_format = 'R$ #,##0.00'

            # Col F: Nº Pessoas (valor já escrito pelo DataFrame, só formata)
            ws[f"F{i}"].alignment = Alignment(horizontal="center")

            # Col G: Valor por Pessoa = Subtotal / Nº Pessoas (com guard /0)
            ws[f"G{i}"] = f'=IF(F{i}=0,"",D{i}/F{i})'
            ws[f"G{i}"].number_format = 'R$ #,##0.00'

            # Formata valores monetários
            ws[f"C{i}"].number_format = 'R$ #,##0.00'

            # Col B alinhado ao centro
            ws[f"B{i}"].alignment = Alignment(horizontal="center")

            ws.row_dimensions[i].height = 16

        aplicar_borda_range(ws, 2, len(itens_enriquecidos) + 2, 1, 7)

        # Linha de total
        linha_total = len(itens_enriquecidos) + 3
        ws.merge_cells(f"A{linha_total}:C{linha_total}")
        ws[f"A{linha_total}"].value = "TOTAL GERAL DA NOTA"
        ws[f"A{linha_total}"].font = Font(bold=True, color=branco)
        ws[f"A{linha_total}"].fill = PatternFill("solid", fgColor=azul_escuro)
        ws[f"A{linha_total}"].alignment = Alignment(horizontal="right")

        ws[f"D{linha_total}"] = f"=SUM(D3:D{linha_total - 1})"
        ws[f"D{linha_total}"].number_format = 'R$ #,##0.00'
        ws[f"D{linha_total}"].font = Font(bold=True, color=branco)
        ws[f"D{linha_total}"].fill = PatternFill("solid", fgColor=azul_escuro)
        ws[f"D{linha_total}"].alignment = Alignment(horizontal="center")
        ws.row_dimensions[linha_total].height = 20
        aplicar_borda_range(ws, linha_total, linha_total, 1, 7)

        # Largura das colunas
        larguras = [42, 12, 22, 18, 20, 14, 22]
        for i, larg in enumerate(larguras, start=1):
            ws.column_dimensions[get_column_letter(i)].width = larg

        # ── Aba Resumo ───────────────────────────────────────────────
        ws_r = writer.book.create_sheet("Resumo")

        ws_r.merge_cells("A1:C2")
        ws_r["A1"].value = "📊  RESUMO DA DIVISÃO"
        ws_r["A1"].font = Font(bold=True, size=14, color=branco)
        ws_r["A1"].fill = PatternFill("solid", fgColor=azul_escuro)
        ws_r["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws_r.row_dimensions[1].height = 20
        ws_r.row_dimensions[2].height = 18

        # Cabeçalho da tabela de resumo
        for col_idx, cab in enumerate(["Pessoa", "Nº de Itens", "Valor Total (R$)"], start=1):
            c = ws_r.cell(row=4, column=col_idx, value=cab)
            c.font = Font(bold=True, color=branco, size=11)
            c.fill = PatternFill("solid", fgColor=azul_medio)
            c.alignment = Alignment(horizontal="center")
        ws_r.row_dimensions[4].height = 22

        n_itens = len(itens_enriquecidos)
        valor_por_pessoa = round(total_geral / n_pessoas, 2)

        for row_idx, pessoa in enumerate(pessoas, start=5):
            fill = PatternFill("solid", fgColor=cinza_linha if row_idx % 2 == 0 else branco)
            ws_r.cell(row=row_idx, column=1, value=pessoa).fill = fill
            ws_r.cell(row=row_idx, column=2, value=n_itens).fill = fill
            c = ws_r.cell(row=row_idx, column=3, value=valor_por_pessoa)
            c.fill = fill
            c.number_format = 'R$ #,##0.00'
            c.font = Font(bold=True, color=verde)
            ws_r.row_dimensions[row_idx].height = 18

        # Total
        linha_tot_r = 5 + n_pessoas
        ws_r.merge_cells(f"A{linha_tot_r}:B{linha_tot_r}")
        ws_r[f"A{linha_tot_r}"].value = "TOTAL GERAL DA NOTA"
        ws_r[f"A{linha_tot_r}"].font = Font(bold=True, color=branco)
        ws_r[f"A{linha_tot_r}"].fill = PatternFill("solid", fgColor=azul_escuro)
        ws_r[f"A{linha_tot_r}"].alignment = Alignment(horizontal="right")
        ws_r[f"C{linha_tot_r}"] = total_geral
        ws_r[f"C{linha_tot_r}"].number_format = 'R$ #,##0.00'
        ws_r[f"C{linha_tot_r}"].font = Font(bold=True, color=branco)
        ws_r[f"C{linha_tot_r}"].fill = PatternFill("solid", fgColor=azul_escuro)
        ws_r[f"C{linha_tot_r}"].alignment = Alignment(horizontal="center")
        ws_r.row_dimensions[linha_tot_r].height = 22

        aplicar_borda_range(ws_r, 4, linha_tot_r, 1, 3)

        for i, larg in enumerate([28, 14, 22], start=1):
            ws_r.column_dimensions[get_column_letter(i)].width = larg

    return caminho
