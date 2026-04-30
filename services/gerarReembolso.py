import pandas as pd
import os
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side
)
from openpyxl.utils import get_column_letter


def gerar_excel_reembolso(
    itens: list[dict],
    solicitante: str,
    cargo: str,
    centro_custo: str,
    justificativa: str,
    data_compra: str,
) -> str:
    """
    Gera um Excel de reembolso empresarial estruturado com cabeçalho e
    tabela de itens. Retorna o caminho do arquivo.
    """
    caminho = "reembolso_empresarial.xlsx"

    # ── Paleta de cores ──────────────────────────────────────────────
    azul_escuro   = "1E3A5F"
    azul_medio    = "1E6FD9"
    azul_claro    = "D6E4F7"
    cinza_linha   = "F2F6FC"
    branco        = "FFFFFF"

    borda_fina = Border(
        left=Side(style="thin", color="B0C4DE"),
        right=Side(style="thin", color="B0C4DE"),
        top=Side(style="thin", color="B0C4DE"),
        bottom=Side(style="thin", color="B0C4DE"),
    )

    def aplicar_borda(ws, min_row, max_row, min_col, max_col):
        for row in ws.iter_rows(
            min_row=min_row, max_row=max_row,
            min_col=min_col, max_col=max_col
        ):
            for cell in row:
                cell.border = borda_fina

    df = pd.DataFrame(itens)
    total_geral = df["Subtotal"].sum()

    with pd.ExcelWriter(caminho, engine="openpyxl") as writer:
        # ── Aba 1: Solicitação ───────────────────────────────────────
        ws = writer.book.create_sheet("Solicitação", 0)
        writer.book.active = ws

        # Cabeçalho institucional
        ws.merge_cells("A1:F2")
        titulo_cell = ws["A1"]
        titulo_cell.value = "📄  SOLICITAÇÃO DE REEMBOLSO"
        titulo_cell.font = Font(bold=True, size=16, color=branco)
        titulo_cell.fill = PatternFill("solid", fgColor=azul_escuro)
        titulo_cell.alignment = Alignment(horizontal="center", vertical="center")

        ws.row_dimensions[1].height = 20
        ws.row_dimensions[2].height = 18

        # Dados do solicitante
        campos = [
            ("Solicitante:", solicitante),
            ("Cargo / Função:", cargo),
            ("Centro de Custo:", centro_custo),
            ("Data da Compra:", data_compra),
            ("Justificativa:", justificativa),
        ]

        for i, (campo, valor) in enumerate(campos, start=4):
            ws[f"A{i}"] = campo
            ws[f"A{i}"].font = Font(bold=True, color=azul_escuro)
            ws[f"A{i}"].fill = PatternFill("solid", fgColor=azul_claro)
            ws.merge_cells(f"B{i}:F{i}")
            ws[f"B{i}"] = valor
            ws[f"B{i}"].alignment = Alignment(wrap_text=True)
            aplicar_borda(ws, i, i, 1, 6)
            ws.row_dimensions[i].height = 18

        # Tabela de itens
        inicio_tabela = 4 + len(campos) + 2
        cabecalhos = ["Produto", "Qtd", "Valor Unitário (R$)", "Subtotal (R$)"]

        for col_idx, cab in enumerate(cabecalhos, start=1):
            c = ws.cell(row=inicio_tabela, column=col_idx, value=cab)
            c.font = Font(bold=True, color=branco, size=11)
            c.fill = PatternFill("solid", fgColor=azul_medio)
            c.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[inicio_tabela].height = 22

        for row_idx, item in enumerate(itens, start=inicio_tabela + 1):
            fill = PatternFill("solid", fgColor=cinza_linha if row_idx % 2 == 0 else branco)
            valores = [
                item["Produto"],
                item["Quantidade"],
                item["Valor Unitário"],
                item["Subtotal"],
            ]
            for col_idx, val in enumerate(valores, start=1):
                c = ws.cell(row=row_idx, column=col_idx, value=val)
                c.fill = fill
                if col_idx in (2, 3, 4):
                    c.alignment = Alignment(horizontal="center")
                if col_idx in (3, 4):
                    c.number_format = 'R$ #,##0.00'
            ws.row_dimensions[row_idx].height = 16

        aplicar_borda(ws, inicio_tabela, inicio_tabela + len(itens), 1, 4)

        # Linha de total
        linha_total = inicio_tabela + len(itens) + 1
        ws.merge_cells(f"A{linha_total}:C{linha_total}")
        ws[f"A{linha_total}"] = "TOTAL GERAL"
        ws[f"A{linha_total}"].font = Font(bold=True, color=branco)
        ws[f"A{linha_total}"].fill = PatternFill("solid", fgColor=azul_escuro)
        ws[f"A{linha_total}"].alignment = Alignment(horizontal="right")
        ws[f"D{linha_total}"] = total_geral
        ws[f"D{linha_total}"].number_format = 'R$ #,##0.00'
        ws[f"D{linha_total}"].font = Font(bold=True, color=branco)
        ws[f"D{linha_total}"].fill = PatternFill("solid", fgColor=azul_escuro)
        ws[f"D{linha_total}"].alignment = Alignment(horizontal="center")
        ws.row_dimensions[linha_total].height = 20
        aplicar_borda(ws, linha_total, linha_total, 1, 4)

        # Largura das colunas
        larguras = [40, 8, 20, 18, 5, 5]
        for i, larg in enumerate(larguras, start=1):
            ws.column_dimensions[get_column_letter(i)].width = larg

        # ── Aba 2: Aprovação ─────────────────────────────────────────
        ws2 = writer.book.create_sheet("Aprovação")

        ws2.merge_cells("A1:D2")
        ws2["A1"].value = "✅  APROVAÇÃO / VISTO"
        ws2["A1"].font = Font(bold=True, size=14, color=branco)
        ws2["A1"].fill = PatternFill("solid", fgColor=azul_escuro)
        ws2["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws2.row_dimensions[1].height = 20
        ws2.row_dimensions[2].height = 18

        campos_aprov = [
            ("Solicitante:", ""),
            ("Gestor direto:", ""),
            ("Financeiro:", ""),
            ("Diretoria:", ""),
            ("Data da aprovação:", ""),
            ("Observações:", ""),
        ]
        for i, (campo, _) in enumerate(campos_aprov, start=4):
            ws2[f"A{i}"] = campo
            ws2[f"A{i}"].font = Font(bold=True, color=azul_escuro)
            ws2[f"A{i}"].fill = PatternFill("solid", fgColor=azul_claro)
            ws2.merge_cells(f"B{i}:D{i}")
            ws2[f"B{i}"] = "___________________________"
            ws2[f"B{i}"].alignment = Alignment(horizontal="center")
            aplicar_borda(ws2, i, i, 1, 4)
            ws2.row_dimensions[i].height = 22

        for i in [1, 2, 3, 4]:
            ws2.column_dimensions[get_column_letter(i)].width = 28

    return caminho
